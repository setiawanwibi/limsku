import openpyxl, json, html

wb = openpyxl.load_workbook(r'd:\xampp\htdocs\limsku\template_kualitatif.xlsx', data_only=True)
ws = wb['Rapih']

def b_str(b):
    if not b: return ''
    p = []
    if b.top and b.top.style: p.append(f'T:{b.top.style}')
    if b.bottom and b.bottom.style: p.append(f'B:{b.bottom.style}')
    if b.left and b.left.style: p.append(f'L:{b.left.style}')
    if b.right and b.right.style: p.append(f'R:{b.right.style}')
    return ','.join(p)

merged_map = {}
for mr in sorted(list(ws.merged_cells.ranges), key=lambda x: (x.min_row, x.min_col)):
    for r in range(mr.min_row, mr.max_row + 1):
        for c in range(mr.min_col, mr.max_col + 1):
            is_origin = (r == mr.min_row and c == mr.min_col)
            merged_map[(r, c)] = {
                'min_r': mr.min_row,
                'max_r': mr.max_row,
                'min_c': mr.min_col,
                'max_c': mr.max_col,
                'rowspan': mr.max_row - mr.min_row + 1,
                'colspan': mr.max_col - mr.min_col + 1,
                'is_origin': is_origin
            }

def get_block_borders(r, c):
    if (r, c) in merged_map:
        m = merged_map[(r, c)]
        if not m['is_origin']:
            return None
        min_r, max_r = m['min_r'], m['max_r']
        min_c, max_c = m['min_c'], m['max_c']
    else:
        min_r, max_r = r, r
        min_c, max_c = c, c

    top = None
    for col in range(min_c, max_c + 1):
        b = ws.cell(min_r, col).border
        if b and b.top and b.top.style:
            top = b.top.style
            break

    bottom = None
    for col in range(min_c, max_c + 1):
        b = ws.cell(max_r, col).border
        if b and b.bottom and b.bottom.style:
            bottom = b.bottom.style
            break

    left = None
    for row in range(min_r, max_r + 1):
        b = ws.cell(row, min_c).border
        if b and b.left and b.left.style:
            left = b.left.style
            break

    right = None
    for row in range(min_r, max_r + 1):
        b = ws.cell(row, max_c).border
        if b and b.right and b.right.style:
            right = b.right.style
            break

    return {'top': top, 'bottom': bottom, 'left': left, 'right': right}

# Let's inspect rows 1 to 95 and write the template code
print("Building view rows...")
