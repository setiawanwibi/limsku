import openpyxl, json, html, re

wb = openpyxl.load_workbook(r'd:\xampp\htdocs\limsku\template_kualitatif.xlsx', data_only=True)
ws = wb['Rapih']

def b_str(b):
    if not b: return ''
    p = []
    if b.top and b.top.style: p.append(f'T:{b.top.style}')
    if b.bottom and b.bottom.style: p.append(f'B:{b.bottom.style}')
    if b.left and b.left.style: p.append(f'L:{b.left.style}')
    if b.right and b.right.style: p.append(f'R:{b.right.style}')
    return ','.join(p)

merged_map = {}
for mr in sorted(list(ws.merged_cells.ranges), key=lambda x: (x.min_row, x.min_col)):
    for r in range(mr.min_row, mr.max_row + 1):
        for c in range(mr.min_col, mr.max_col + 1):
            is_origin = (r == mr.min_row and c == mr.min_col)
            merged_map[(r, c)] = {
                'min_r': mr.min_row,
                'max_r': mr.max_row,
                'min_c': mr.min_col,
                'max_c': mr.max_col,
                'rowspan': mr.max_row - mr.min_row + 1,
                'colspan': mr.max_col - mr.min_col + 1,
                'is_origin': is_origin
            }

def get_block_borders(r, c):
    if (r, c) in merged_map:
        m = merged_map[(r, c)]
        if not m['is_origin']:
            return None
        min_r, max_r = m['min_r'], m['max_r']
        min_c, max_c = m['min_c'], m['max_c']
    else:
        min_r, max_r = r, r
        min_c, max_c = c, c

    top = None
    for col in range(min_c, max_c + 1):
        b = ws.cell(min_r, col).border
        if b and b.top and b.top.style:
            top = b.top.style
            break

    bottom = None
    for col in range(min_c, max_c + 1):
        b = ws.cell(max_r, col).border
        if b and b.bottom and b.bottom.style:
            bottom = b.bottom.style
            break

    left = None
    for row in range(min_r, max_r + 1):
        b = ws.cell(row, min_c).border
        if b and b.left and b.left.style:
            left = b.left.style
            break

    right = None
    for row in range(min_r, max_r + 1):
        b = ws.cell(row, max_c).border
        if b and b.right and b.right.style:
            right = b.right.style
            break

    # Also outer border enforcement: row 1 top is medium, rows 1-94 left is medium, right is medium, row 94 bottom is medium
    if min_r == 1: top = 'medium'
    if max_r == 94: bottom = 'medium'
    if min_c == 1: left = 'medium'
    if max_c == 20: right = 'medium'

    return {'t': top, 'b': bottom, 'l': left, 'r': right}

col_widths_char = [ws.column_dimensions[col_let].width for col_let in 'ABCDEFGHIJKLMNOPQRST']
tot_chars = sum(col_widths_char)
col_widths_pct = [round((ch / tot_chars) * 100, 3) for ch in col_widths_char]

print(f"Total columns: {len(col_widths_pct)}, sum pct: {sum(col_widths_pct)}")

# Save to json helper
spec_rows = []
for r in range(1, 96):
    rd = ws.row_dimensions.get(r)
    h = rd.height if rd else 15.0
    cells = []
    for c in range(1, 21):
        if (r, c) in merged_map and not merged_map[(r, c)]['is_origin']:
            continue
        cell = ws.cell(r, c)
        val = cell.value
        f = cell.font
        fn = f.name if f else 'Times New Roman'
        fs = f.size if f else 11.0
        bold = bool(f.bold) if f else False
        italic = bool(f.italic) if f else False
        al = cell.alignment
        ah = al.horizontal if (al and al.horizontal) else 'left'
        av = al.vertical if (al and al.vertical) else 'middle'
        wrap = bool(al.wrap_text) if al else False
        borders = get_block_borders(r, c)
        
        rowspan = merged_map[(r, c)]['rowspan'] if (r, c) in merged_map else 1
        colspan = merged_map[(r, c)]['colspan'] if (r, c) in merged_map else 1

        cells.append({
            'c': c,
            'rs': rowspan,
            'cs': colspan,
            'v': val,
            'fn': fn,
            'fs': fs,
            'b': bold,
            'i': italic,
            'ah': ah,
            'av': av,
            'wrap': wrap,
            'borders': borders
        })
    spec_rows.append({
        'r': r,
        'h': h,
        'cells': cells
    })

with open('template_processed_grid.json', 'w', encoding='utf-8') as f:
    json.dump({'col_pct': col_widths_pct, 'rows': spec_rows}, f, ensure_ascii=False, indent=2)

print("Saved template_processed_grid.json successfully!")
