import openpyxl, json, html, re

wb = openpyxl.load_workbook(r'd:\xampp\htdocs\limsku\template_kualitatif.xlsx', data_only=True)
ws = wb['Rapih']

def b_str(b):
    if not b: return ''
    p = []
    if b.top and b.top.style: p.append(f'T:{b.top.style}')
    if b.bottom and b.bottom.style: p.append(f'B:{b.bottom.style}')
    if b.left and b.left.style: p.append(f'L:{b.left.style}')
    if b.right and b.right.style: p.append(f'R:{b.right.style}')
    return ','.join(p)

merged_map = {}
for mr in sorted(list(ws.merged_cells.ranges), key=lambda x: (x.min_row, x.min_col)):
    for r in range(mr.min_row, mr.max_row + 1):
        for c in range(mr.min_col, mr.max_col + 1):
            is_origin = (r == mr.min_row and c == mr.min_col)
            merged_map[(r, c)] = {
                'min_r': mr.min_row,
                'max_r': mr.max_row,
                'min_c': mr.min_col,
                'max_c': mr.max_col,
                'rowspan': mr.max_row - mr.min_row + 1,
                'colspan': mr.max_col - mr.min_col + 1,
                'is_origin': is_origin
            }

def get_block_borders(r, c):
    if (r, c) in merged_map:
        m = merged_map[(r, c)]
        if not m['is_origin']:
            return None
        min_r, max_r = m['min_r'], m['max_r']
        min_c, max_c = m['min_c'], m['max_c']
    else:
        min_r, max_r = r, r
        min_c, max_c = c, c

    top = None
    for col in range(min_c, max_c + 1):
        b = ws.cell(min_r, col).border
        if b and b.top and b.top.style:
            top = b.top.style
            break

    bottom = None
    for col in range(min_c, max_c + 1):
        b = ws.cell(max_r, col).border
        if b and b.bottom and b.bottom.style:
            bottom = b.bottom.style
            break

    left = None
    for row in range(min_r, max_r + 1):
        b = ws.cell(row, min_c).border
        if b and b.left and b.left.style:
            left = b.left.style
            break

    right = None
    for row in range(min_r, max_r + 1):
        b = ws.cell(row, max_c).border
        if b and b.right and b.right.style:
            right = b.right.style
            break

    # Outer border enforcement
    if min_r == 1: top = 'medium'
    if max_r == 94: bottom = 'medium'
    if min_c == 1: left = 'medium'
    if max_c == 20: right = 'medium'

    return {'t': top, 'b': bottom, 'l': left, 'r': right}

col_widths_char = [ws.column_dimensions[col_let].width for col_let in 'ABCDEFGHIJKLMNOPQRST']
tot_chars = sum(col_widths_char)
col_widths_pct = [round((ch / tot_chars) * 100, 3) for ch in col_widths_char]

# Map of dynamic cell substitutions:
# key: (r, c)
# value: PHP snippet string to render inside cell
dynamic_replacements = {
    # Target mikroba (R5-R7)
    (5, 5): "<?php echo (!empty($p['target_staphylococcus']) ? '&#9745;' : '&#9744;'); ?>",
    (5, 10): "<?php echo (!empty($p['target_candida']) ? '&#9745;' : '&#9744;'); ?>",
    (5, 14): "<?php echo (!empty($p['target_salmonella']) ? '&#9745;' : '&#9744;'); ?>",
    (6, 5): "<?php echo (!empty($p['target_pseudomonas']) ? '&#9745;' : '&#9744;'); ?>",
    (6, 10): "<?php echo (!empty($p['target_shigella']) ? '&#9745;' : '&#9744;'); ?>",
    (6, 14): "<?php echo (!empty($p['target_ecoli']) ? '&#9745;' : '&#9744;'); ?>",
    (7, 14): "<?php echo (!empty($p['target_lainnya_check']) ? '&#9745;' : '&#9744;'); ?>",
    (7, 15): "Lainnya: <?php echo !empty($p['target_lainnya_text']) ? '<u>' . html_escape($p['target_lainnya_text']) . '</u>' : '..............'; ?>",

    # Lab location (R8)
    (8, 2): "Laboratorium Mikrobiologi Balai Besar POM di Bandar Lampung",

    # Header info sampel (R10-R12)
    (10, 2): "Nama Sampel Uji : <strong style=\"font-size: 11pt; color: #000;\"><?php echo html_escape($sampel['nama_sampel'] ?? '-'); ?></strong>",
    (10, 9): "Kode Sampel Uji : <strong style=\"font-size: 11pt; color: #000;\"><?php echo html_escape(!empty($sampel['kode_sampel_manual']) ? $sampel['kode_sampel_manual'] : (!empty($sampel['no']) ? 'NO-' . $sampel['no'] : 'SMP-' . $sampel['id'])); ?></strong>",
    (10, 18): "<span style=\"font-size: 10pt;\"><?php echo !empty($sampel['tanggal_sampling']) ? date('d-m-Y', strtotime($sampel['tanggal_sampling'])) : (!empty($sampel['created_at']) ? date('d-m-Y', strtotime($sampel['created_at'])) : '-'); ?></span>",
    (11, 18): "<span style=\"font-size: 10pt; font-weight: bold;\"><?php echo html_escape($penguji['nama_lengkap'] ?? ($sesi['nama_penguji'] ?? '-')); ?></span>",
    (12, 18): "<span style=\"font-size: 10pt;\"><?php echo !empty($sesi['waktu_mulai']) ? date('d-m-Y', strtotime($sesi['waktu_mulai'])) : date('d-m-Y'); ?></span>",

    # Kesesuaian Metode (R15-R19)
    (15, 5): "<?php echo (isset($p['pelarut_kesesuaian']) && $p['pelarut_kesesuaian'] === 'MLB') ? '&#9745;' : '&#9744;'; ?>",
    (16, 5): "<?php echo (isset($p['pelarut_kesesuaian']) && $p['pelarut_kesesuaian'] === 'TSB') ? '&#9745;' : '&#9744;'; ?>",
    (17, 5): "<?php echo (isset($p['pelarut_kesesuaian']) && $p['pelarut_kesesuaian'] === 'TSB + 30 g/L polysorbate 80') ? '&#9745;' : '&#9744;'; ?>",
    (18, 5): "<?php echo (isset($p['pelarut_kesesuaian']) && $p['pelarut_kesesuaian'] === 'TSB + 30 g/L polysorbate 80 + 3 g/L Lesitin') ? '&#9745;' : '&#9744;'; ?>",
    (19, 5): "<?php echo (isset($p['pelarut_kesesuaian']) && $p['pelarut_kesesuaian'] === 'Lainnya') ? '&#9745;' : '&#9744;'; ?>",
    (19, 6): "Lainnya: <?php echo !empty($p['pelarut_kesesuaian_lainnya']) ? '<u>' . html_escape($p['pelarut_kesesuaian_lainnya']) . '</u>' : '…..........................'; ?>",

    (16, 12): "<?php echo (!empty($p['baku_staphylococcus']) ? '&#9745;' : '&#9744;'); ?>",
    (16, 13): "Staphylococcus aures ATCC / WDCM <?php echo !empty($p['baku_staphylococcus_no']) ? '<u>' . html_escape($p['baku_staphylococcus_no']) . '</u>' : '…................'; ?>",
    (17, 12): "<?php echo (!empty($p['baku_pseudomonas']) ? '&#9745;' : '&#9744;'); ?>",
    (17, 13): "Pseudomonas aeruginosa ATCC / WDCM <?php echo !empty($p['baku_pseudomonas_no']) ? '<u>' . html_escape($p['baku_pseudomonas_no']) . '</u>' : '…...............'; ?>",
    (18, 12): "<?php echo (!empty($p['baku_candida']) ? '&#9745;' : '&#9744;'); ?>",
    (18, 13): "Candida albicans ATCC / WDCM <?php echo !empty($p['baku_candida_no']) ? '<u>' . html_escape($p['baku_candida_no']) . '</u>' : '…...............'; ?>",
    (19, 12): "<?php echo (!empty($p['baku_lainnya']) ? '&#9745;' : '&#9744;'); ?>",
    (19, 13): "Lainnya: <?php echo !empty($p['baku_lainnya']) ? '<u>' . html_escape($p['baku_lainnya']) . '</u>' : '…..........................'; ?>",

    # Homogenisasi awal (R22-R27)
    (22, 6): "<strong style=\"font-size: 11pt;\"><?php echo html_escape($p['bobot_sampel'] ?? '10.0'); ?></strong>",
    (23, 6): "<strong style=\"font-size: 11pt;\"><?php echo html_escape($p['volume_pelarut'] ?? '90.0'); ?></strong>",
    (23, 8): "<?php echo (isset($p['jenis_pelarut_homogenisasi']) && $p['jenis_pelarut_homogenisasi'] === 'MLB') ? '&#9745;' : '&#9744;'; ?>",
    (24, 8): "<?php echo (isset($p['jenis_pelarut_homogenisasi']) && $p['jenis_pelarut_homogenisasi'] === 'TSB') ? '&#9745;' : '&#9744;'; ?>",
    (25, 8): "<?php echo (isset($p['jenis_pelarut_homogenisasi']) && $p['jenis_pelarut_homogenisasi'] === 'TSB + 30 g/L polysorbate 80') ? '&#9745;' : '&#9744;'; ?>",
    (26, 8): "<?php echo (isset($p['jenis_pelarut_homogenisasi']) && $p['jenis_pelarut_homogenisasi'] === 'TSB + 30 g/L polysorbate 80 + 3 g/L Lesitin') ? '&#9745;' : '&#9744;'; ?>",
    (27, 8): "<?php echo (isset($p['jenis_pelarut_homogenisasi']) && $p['jenis_pelarut_homogenisasi'] === 'Lainnya') ? '&#9745;' : '&#9744;'; ?>",
    (27, 9): "Lainnya : <?php echo !empty($p['pelarut_homogenisasi_lainnya']) ? '<u>' . html_escape($p['pelarut_homogenisasi_lainnya']) . '</u>' : '….......................'; ?>",

    # Inokulum tingkat (R31-R33)
    (31, 12): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === '1 g atau mL') ? '&#9745;' : '&#9744;'; ?>",
    (31, 14): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === '1 g atau mL') ? '&#9745;' : '&#9744;'; ?>",
    (31, 16): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === '1 g atau mL') ? '&#9745;' : '&#9744;'; ?>",
    (31, 18): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === '1 g atau mL') ? '&#9745;' : '&#9744;'; ?>",

    (32, 12): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === '0,1 g atau mL') ? '&#9745;' : '&#9744;'; ?>",
    (32, 14): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === '0,1 g atau mL') ? '&#9745;' : '&#9744;'; ?>",
    (32, 16): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === '0,1 g atau mL') ? '&#9745;' : '&#9744;'; ?>",
    (32, 18): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === '0,1 g atau mL') ? '&#9745;' : '&#9744;'; ?>",

    (33, 12): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === 'Lainnya') ? '&#9745;' : '&#9744;'; ?>",
    (33, 14): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === 'Lainnya') ? '&#9745;' : '&#9744;'; ?>",
    (33, 16): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === 'Lainnya') ? '&#9745;' : '&#9744;'; ?>",
    (33, 18): "<?php echo (isset($p['inokulum_tingkat']) && $p['inokulum_tingkat'] === 'Lainnya') ? '&#9745;' : '&#9744;'; ?>",

    # Stage 0: Pra Pengayaan (R34)
    (34, 2): "<?php echo !empty($t0['tanggal']) ? date('d-m-Y', strtotime($t0['tanggal'])) : (!empty($contoh_mode) ? 'N/A' : ''); ?>",
    (34, 12): "<?php echo !empty($t0['sampel_spike']) ? html_escape($t0['sampel_spike']) : (!empty($contoh_mode) ? 'N/A' : ''); ?>",
    (34, 14): "<?php echo !empty($t0['media_spike']) ? html_escape($t0['media_spike']) : (!empty($contoh_mode) ? 'N/A' : ''); ?>",
    (34, 16): "<?php echo !empty($t0['kontrol_media']) ? html_escape($t0['kontrol_media']) : (!empty($contoh_mode) ? 'N/A' : ''); ?>",
    (34, 18): "<?php echo !empty($t0['sampel_uji']) ? html_escape($t0['sampel_uji']) : (!empty($contoh_mode) ? 'N/A' : ''); ?>",

    # Stage 1: Pengayaan (R37)
    (37, 2): "<?php echo !empty($t1['tanggal']) ? date('d-m-Y', strtotime($t1['tanggal'])) : ''; ?>",
    (37, 12): "<?php echo !empty($t1['sampel_spike']) ? html_escape($t1['sampel_spike']) : ''; ?>",
    (37, 14): "<?php echo !empty($t1['media_spike']) ? html_escape($t1['media_spike']) : 'Keruh'; ?>",
    (37, 16): "<?php echo !empty($t1['kontrol_media']) ? html_escape($t1['kontrol_media']) : ''; ?>",
    (37, 18): "<?php echo !empty($t1['sampel_uji']) ? html_escape($t1['sampel_uji']) : ''; ?>",

    # Stage 2: Seleksi (R45)
    (45, 2): "<?php echo !empty($t2['tanggal']) ? date('d-m-Y', strtotime($t2['tanggal'])) : ''; ?>",
    (45, 12): "<?php echo !empty($t2['sampel_spike']) ? html_escape($t2['sampel_spike']) : ''; ?>",
    (45, 14): "<?php echo !empty($t2['media_spike']) ? html_escape($t2['media_spike']) : 'Keruh'; ?>",
    (45, 16): "<?php echo !empty($t2['kontrol_media']) ? html_escape($t2['kontrol_media']) : ''; ?>",
    (45, 18): "<?php echo !empty($t2['sampel_uji']) ? html_escape($t2['sampel_uji']) : ''; ?>",

    # Stage 3: Isolasi (R50)
    (50, 2): "<?php echo !empty($t3['tanggal']) ? date('d-m-Y', strtotime($t3['tanggal'])) : ''; ?>",
    (50, 12): "<?php echo !empty($t3['sampel_spike']) ? html_escape($t3['sampel_spike']) : ''; ?>",
    (50, 14): "<?php echo !empty($t3['media_spike']) ? html_escape($t3['media_spike']) : 'Positif'; ?>",
    (50, 16): "<?php echo !empty($t3['kontrol_media']) ? html_escape($t3['kontrol_media']) : ''; ?>",
    (50, 18): "<?php echo !empty($t3['sampel_uji']) ? html_escape($t3['sampel_uji']) : ''; ?>",

    # Stage 4: Konfirmasi (R60)
    (60, 2): "<?php echo !empty($t4['tanggal']) ? date('d-m-Y', strtotime($t4['tanggal'])) : ''; ?>",
    (60, 12): "<?php echo !empty($t4['sampel_spike']) ? html_escape($t4['sampel_spike']) : ''; ?>",
    (60, 14): "<?php echo !empty($t4['media_spike']) ? nl2br(html_escape($t4['media_spike'])) : \"Biokimia Kit / Konvensional *)<br><br>...............................<br><br>\"; ?>",
    (60, 16): "<?php echo !empty($t4['kontrol_media']) ? html_escape($t4['kontrol_media']) : 'N/A'; ?>",
    (60, 18): "<?php echo !empty($t4['sampel_uji']) ? html_escape($t4['sampel_uji']) : ''; ?>",

    # Hasil Uji, Pustaka, Syarat, Kesimpulan (R84-R89)
    (84, 4): ": <strong style=\"font-size: 11pt; color: #000;\"><?php echo html_escape($p['hasil_uji'] ?? ($hasil['hasil_uji'] ?? '-')); ?></strong>",
    (86, 4): ": <span style=\"font-size: 10.5pt;\"><?php echo html_escape($p['pustaka_metode'] ?? ($hasil['pustaka_acuan'] ?? ($metode['pustaka_acuan'] ?? '-'))); ?></span>",
    (86, 14): "Syarat : <strong style=\"font-size: 10.5pt; color: #c00;\"><?php echo html_escape($p['syarat'] ?? ($hasil['persyaratan'] ?? ($metode['persyaratan'] ?? '-'))); ?></strong>",
    (89, 4): """<?php
        $ks = $p['kesimpulan_lcp'] ?? ($hasil['kesimpulan'] ?? '');
        $ms_style = (strpos($ks, 'MS') !== false && strpos($ks, 'TMS') === false) ? 'font-weight: bold; background: #e8f5e9; padding: 2px 6px; border: 1.5px solid #2e7d32;' : 'color: #333;';
        $tms_style = (strpos($ks, 'TMS') !== false) ? 'font-weight: bold; background: #ffebee; padding: 2px 6px; border: 1.5px solid #c62828;' : 'color: #333;';
        $hpst_style = (strpos($ks, 'HPST') !== false) ? 'font-weight: bold; background: #e3f2fd; padding: 2px 6px; border: 1.5px solid #1565c0;' : 'color: #333;';
    ?>
    <span style="<?php echo $ms_style; ?>">Memenuhi Syarat (MS)</span> / 
    <span style="<?php echo $tms_style; ?>">Tidak Memenuhi Syarat (TMS)</span> / 
    <span style="<?php echo $hpst_style; ?>">Hasil Pengujian Seperti Tersebut (HPST)</span> *)
    """,

    # Signatures (R91-R93)
    (91, 2): "Tanda Tangan Penguji :",
    (91, 14): "  Diperiksa oleh :",
    (92, 2): """<?php if (!empty($ttd_penguji_img)): ?>
        <div style="height: 48px; display: flex; align-items: center; justify-content: flex-start; padding-left: 20px;">
            <img src="<?php echo $ttd_penguji_img; ?>" style="max-height: 44px; max-width: 140px; object-fit: contain;" alt="TTD Penguji">
        </div>
    <?php else: ?>
        <div style="height: 48px; display: flex; align-items: center; padding-left: 20px; color: #777; font-size: 9pt; font-style: italic;">
            [ Tanda Tangan Penguji ]
        </div>
    <?php endif; ?>
    <div style=\"font-size: 9.5pt; font-weight: bold; padding-left: 20px;\"><?php echo html_escape($penguji['nama_lengkap'] ?? ($sesi['nama_penguji'] ?? '-')); ?></div>
    """,
    (92, 14): """<?php if ($status_sesi === 'Approved / Final' || $status_sesi === 'Menunggu Approval'): ?>
        <?php if (!empty($ttd_penyelia_img)): ?>
            <div style="height: 48px; display: flex; align-items: center; justify-content: flex-start; padding-left: 20px;">
                <img src="<?php echo $ttd_penyelia_img; ?>" style="max-height: 44px; max-width: 140px; object-fit: contain;" alt="TTD Penyelia">
            </div>
        <?php else: ?>
            <div style="height: 48px; display: flex; align-items: center; padding-left: 20px; color: #2e7d32; font-size: 8.5pt; font-weight: bold;">
                [ VERIFIED DIGITAL ]
            </div>
        <?php endif; ?>
        <div style=\"font-size: 9.5pt; font-weight: bold; padding-left: 20px;\"><?php echo html_escape($verifier['nama_lengkap'] ?? ($sesi['nama_verifier'] ?? '-')); ?></div>
    <?php else: ?>
        <div style="height: 48px; display: flex; align-items: center; padding-left: 20px; color: #999; font-size: 8.5pt; font-style: italic;">
            (Menunggu Verifikasi Penyelia)
        </div>
        <div style=\"font-size: 9.5pt; color: #777; padding-left: 20px;\">(TTD, Nama)</div>
    <?php endif; ?>
    """,
    (93, 2): "Tanggal : <span style=\"font-size: 10pt;\"><?php echo !empty($sesi['waktu_selesai']) ? date('d-m-Y', strtotime($sesi['waktu_selesai'])) : date('d-m-Y'); ?></span>",
    (93, 14): "  Tanggal : <span style=\"font-size: 10pt;\"><?php echo ($status_sesi === 'Approved / Final' || $status_sesi === 'Menunggu Approval') && !empty($sesi['waktu_verifikasi']) ? date('d-m-Y', strtotime($sesi['waktu_verifikasi'])) : '-'; ?></span>"
}

print("Dynamic replacements dictionary ready.")
